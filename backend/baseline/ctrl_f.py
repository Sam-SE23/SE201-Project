"""Ctrl+F baseline: plain keyword search over the raw document text.

This is the conventional, non-AI alternative the project is evaluated against.
It deliberately does not use the AI pipeline (no chunking, no embeddings), so
the comparison stays fair.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

SNIPPET_RADIUS = 80


@dataclass
class Page:
    source_file: str
    location: str             # "p.3" or "slide 7"
    text: str


@dataclass
class Hit:
    source_file: str
    location: str
    snippet: str


def extract_pages(path: Path) -> list[Page]:
    """Extract text per page (PDF) or per slide (PPTX)."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    raise ValueError(f"Unsupported file type: {path.name} (expected .pdf or .pptx)")


def _extract_pdf(path: Path) -> list[Page]:
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(Page(path.name, f"p.{i}", text))
    return pages


def _extract_pptx(path: Path) -> list[Page]:
    from pptx import Presentation

    pages = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = "\n".join(p for p in parts if p.strip())
        if text.strip():
            pages.append(Page(path.name, f"slide {i}", text))
    return pages


class CtrlFBaseline:
    def __init__(self) -> None:
        self._pages: list[Page] = []

    def index(self, paths: list[Path]) -> None:
        indexed = {p.name for p in paths}
        # Re-uploading a file replaces its old pages instead of duplicating them.
        self._pages = [p for p in self._pages if p.source_file not in indexed]
        for path in paths:
            self._pages.extend(extract_pages(path))

    def search(self, keyword: str) -> list[Hit]:
        """Exact, case-insensitive substring match. Returns every occurrence."""
        keyword = keyword.strip()
        if not keyword:
            return []
        needle = keyword.lower()
        hits = []
        for page in self._pages:
            haystack = page.text.lower()
            start = haystack.find(needle)
            while start != -1:
                hits.append(Hit(page.source_file, page.location,
                                _snippet(page.text, start, len(needle))))
                start = haystack.find(needle, start + len(needle))
        return hits


def _snippet(text: str, start: int, length: int) -> str:
    lo = max(0, start - SNIPPET_RADIUS)
    hi = min(len(text), start + length + SNIPPET_RADIUS)
    body = text[lo:hi].replace("\n", " ")
    return ("…" if lo > 0 else "") + body + ("…" if hi < len(text) else "")
