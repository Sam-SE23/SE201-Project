"""
Step 1: Ingestion.

Pulls raw text out of PDFs and PPTX files in DATA_DIR, keeping the
source filename + page/slide number attached to every block of text.
note: This metadata is what lets Member 2 show citations later, so don't drop it.
"""

from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from pptx import Presentation

import sys
sys.path.append(str(Path(__file__).parent.parent))
from config import DATA_DIR


@dataclass
class RawPage:
    """One page/slide of raw text, before chunking."""
    source: str       # filename
    page: int         # 1-indexed page or slide number
    text: str


def load_pdf(path: Path) -> list[RawPage]:
    pages = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                pages.append(RawPage(source=path.name, page=i, text=text))
    return pages


def load_pptx(path: Path) -> list[RawPage]:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        texts.append(line.strip())
            # pick up any speaker notes too — often has useful context
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[notes] {notes}")
        text = "\n".join(texts).strip()
        if text:
            slides.append(RawPage(source=path.name, page=i, text=text))
    return slides


def load_documents(data_dir: Path = DATA_DIR) -> list[RawPage]:
    """Walk DATA_DIR and load every supported file into RawPage objects."""
    data_dir = Path(data_dir)
    all_pages: list[RawPage] = []

    files = sorted(data_dir.glob("*"))
    if not files:
        print(f"[ingest] No files found in {data_dir}. Drop PDFs/PPTX there first.")
        return all_pages

    skipped: list[str] = []

    for path in files:
        if path.is_dir():
            continue  # subfolders are not read
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                pages = load_pdf(path)
            elif suffix == ".pptx":
                pages = load_pptx(path)
            else:
                print(f"[ingest] SKIPPED {path.name}: unsupported file type '{suffix or 'none'}' (only .pdf and .pptx)")
                skipped.append(path.name)
                continue
            print(f"[ingest] {path.name}: {len(pages)} pages/slides with text")
            all_pages.extend(pages)
        except Exception as e:
            # Don't let one bad file kill the whole ingestion run.
            print(f"[ingest] WARNING: failed to parse {path.name}: {e}")
            skipped.append(path.name)

    if skipped:
        print(f"[ingest] {len(skipped)} file(s) not ingested: {', '.join(skipped)}")

    return all_pages


if __name__ == "__main__":
    pages = load_documents()
    print(f"\n[ingest] Total pages/slides loaded: {len(pages)}")
    if pages:
        print("\n--- sample ---")
        print(f"source={pages[0].source} page={pages[0].page}")
        print(pages[0].text[:300])