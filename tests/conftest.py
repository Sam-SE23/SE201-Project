from pathlib import Path

import pytest
from pptx import Presentation
from reportlab.pdfgen import canvas

from backend import registry
from backend.orchestrator import Orchestrator
from backend.stubs import FakeLLM, FakeQA, FakeQuiz, FakeRetriever

PDF_PAGES = [
    ["The project proposal is due on 2 October at 23:59.",
     "The proposal must include a problem statement and an evaluation plan."],
    ["The checkpoint is due on 18 October at 23:59.",
     "Teams must demonstrate one working core feature at the checkpoint."],
]
SLIDES = [
    "Retrieval augmented generation grounds answers in source documents.",
    "Hallucination rate measures answers citing content absent from the sources.",
]


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "syllabus.pdf"
    c = canvas.Canvas(str(path))
    for lines in PDF_PAGES:
        y = 800
        for line in lines:
            c.drawString(50, y, line)
            y -= 20
        c.showPage()
    c.save()
    return path


@pytest.fixture
def sample_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "lecture.pptx"
    prs = Presentation()
    for text in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Lecture"
        slide.placeholders[1].text = text
    prs.save(path)
    return path


@pytest.fixture
def orch(tmp_path: Path) -> Orchestrator:
    components = registry.Components(FakeLLM(), FakeRetriever(), FakeQA(), FakeQuiz(),
                                     status={"llm": "stub", "retriever": "stub",
                                             "qa": "stub", "quiz": "stub"})
    return Orchestrator(components, upload_dir=tmp_path / "uploads",
                        log_path=tmp_path / "logs" / "ai_log.jsonl")
